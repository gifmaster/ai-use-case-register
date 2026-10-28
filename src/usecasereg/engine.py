"""AI use case register: load, validate, calculate governance intensity, check register health.

The schema, the intensity matrix and the health checks live in YAML files in ``rules/``.
The code applies them, so every result can be traced to a rule.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

import yaml

RULES_DIR = Path(__file__).parent / "rules"
SEVERITY_ORDER = {"high": 0, "medium": 1, "low": 2}

Row = dict[str, str]


def _load(name: str) -> dict:
    return yaml.safe_load((RULES_DIR / name).read_text(encoding="utf-8"))


# --------------------------------------------------------------------------- #
# Schema
# --------------------------------------------------------------------------- #


class Schema:
    def __init__(self, data: dict | None = None):
        data = data or _load("schema.yaml")
        self.columns: list[dict] = data["columns"]
        self.ids = [c["id"] for c in self.columns]
        self.labels = {c["id"]: c["label"] for c in self.columns}
        self.allowed = {c["id"]: [str(v) for v in c["allowed"]] for c in self.columns if "allowed" in c}
        self.dates = {c["id"] for c in self.columns if c.get("type") == "date"}
        self.oversight_levels = {str(k): v for k, v in data["oversight_levels"].items()}
        self.intensity = {t: {str(k): v for k, v in m.items()} for t, m in data["intensity"].items()}
        self.review_months = data["review_months"]
        self.last_reviewed: str = str(data.get("last_reviewed", ""))


SCHEMA = Schema()


# --------------------------------------------------------------------------- #
# Load and save
# --------------------------------------------------------------------------- #


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (date,)):
        return value.isoformat()[:10]
    s = str(value).strip()
    return "" if s.lower() in ("nan", "none", "nat") else s


def normalise(rows: list[dict], schema: Schema = SCHEMA) -> list[Row]:
    """Keep schema columns in order, clean values, lower-case controlled vocabularies."""
    out = []
    for r in rows:
        row = {c: _clean(r.get(c, "")) for c in schema.ids}
        for c in schema.allowed:
            row[c] = row[c].lower()
        for c in schema.dates:
            row[c] = row[c][:10]
        if any(row.values()):
            out.append(row)
    return out


def read_csv(data: bytes | str) -> list[Row]:
    text = data.decode("utf-8-sig") if isinstance(data, bytes) else data.lstrip("﻿")
    sample = text[:2000]
    delim = ";" if sample.count(";") > sample.count(",") else ","
    return normalise(list(csv.DictReader(io.StringIO(text), delimiter=delim)))


def read_xlsx(data: bytes) -> list[Row]:
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    it = ws.iter_rows(values_only=True)
    header = [_clean(h) for h in next(it)]
    return normalise([dict(zip(header, r)) for r in it])


def read_any(data: bytes, filename: str) -> list[Row]:
    if filename.lower().endswith(".xlsx"):
        return read_xlsx(data)
    if filename.lower().endswith((".csv", ".txt")):
        return read_csv(data)
    raise ValueError("Use a CSV or Excel (.xlsx) file.")


def missing_columns(header: list[str], schema: Schema = SCHEMA) -> list[str]:
    return [c for c in schema.ids if c not in header]


def to_csv(rows: list[Row], schema: Schema = SCHEMA) -> str:
    out = io.StringIO()
    w = csv.DictWriter(out, fieldnames=schema.ids, lineterminator="\n")
    w.writeheader()
    for r in rows:
        w.writerow({c: r.get(c, "") for c in schema.ids})
    return out.getvalue()


PLATFORM_HEADERS = {
    "use_case_id": "ID", "name": "Name", "purpose": "Description", "business_unit": "Business unit",
    "use_case_owner": "Owner", "ai_system": "AI system", "vendor": "Vendor", "our_role": "Our role",
    "affected_people": "Affected people", "personal_data": "Personal data", "oversight_level": "Human oversight level",
    "eu_ai_act_tier": "EU AI Act risk tier", "governance_intensity": "Governance intensity",
    "dpia_required": "DPIA required", "fria_required": "FRIA required", "status": "Status",
    "approved_by": "Approved by", "approved_on": "Approved on", "next_review": "Next review", "notes": "Notes",
}
YES_NO = {"yes": "Yes", "no": "No"}


def to_platform_csv(rows: list[Row], schema: Schema = SCHEMA) -> str:
    """CSV with plain-language headers and Yes/No values, to ease mapping into a GRC platform's import wizard.

    This is not any vendor's native template. Platforms let you map columns on import; the headers
    here are chosen to be easy to map, and the mapping is always yours to check.
    """
    out = io.StringIO()
    fields = list(schema.ids)
    w = csv.writer(out, lineterminator="\n")
    w.writerow([PLATFORM_HEADERS.get(c, c) for c in fields])
    for r in rows:
        w.writerow([YES_NO.get(str(r.get(c, "")).lower(), r.get(c, "")) if c in ("personal_data", "dpia_required", "fria_required") else r.get(c, "") for c in fields])
    return out.getvalue()


def to_xlsx(rows: list[Row], issues: list["Issue"], schema: Schema = SCHEMA) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    ws.append(schema.ids)
    for r in rows:
        ws.append([r.get(c, "") for c in schema.ids])
    hs = wb.create_sheet("Health check")
    hs.append(["use_case_id", "name", "severity", "issue", "action", "reference"])
    for i in issues:
        hs.append([i.use_case_id, i.name, i.severity, i.title, i.action, i.ref])
    for sheet in (ws, hs):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1A1A2E")
        sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Validation and intensity
# --------------------------------------------------------------------------- #


def value_errors(rows: list[Row], schema: Schema = SCHEMA) -> list[str]:
    errors, seen = [], set()
    for n, r in enumerate(rows, start=1):
        rid = r.get("use_case_id") or f"row {n}"
        if not r.get("use_case_id"):
            errors.append(f"Row {n}: missing use_case_id")
        elif r["use_case_id"] in seen:
            errors.append(f"{rid}: duplicate use_case_id")
        seen.add(r.get("use_case_id"))
        for c, allowed in schema.allowed.items():
            if r.get(c) and r[c] not in allowed:
                errors.append(f"{rid}: {c} = '{r[c]}' is not one of {', '.join(allowed)}")
        for c in schema.dates:
            if r.get(c):
                try:
                    date.fromisoformat(r[c])
                except ValueError:
                    errors.append(f"{rid}: {c} = '{r[c]}' is not a date (YYYY-MM-DD)")
    return errors


def expected_intensity(row: Row, schema: Schema = SCHEMA) -> str | None:
    tier, level = row.get("eu_ai_act_tier"), row.get("oversight_level")
    if tier in schema.intensity and level in schema.intensity[tier]:
        return schema.intensity[tier][level]
    if tier == "prohibited":
        return "stop"
    return None


def suggested_next_review(row: Row, from_date: date, schema: Schema = SCHEMA) -> str | None:
    months = schema.review_months.get(row.get("governance_intensity") or expected_intensity(row) or "")
    if not months:
        return None
    y, m = divmod(from_date.month - 1 + months, 12)
    day = min(from_date.day, 28)
    return date(from_date.year + y, m + 1, day).isoformat()


# --------------------------------------------------------------------------- #
# Health checks
# --------------------------------------------------------------------------- #


@dataclass
class Issue:
    use_case_id: str
    name: str
    check_id: str
    severity: str
    title: str
    action: str
    ref: str


class Checks:
    OPS = {"eq", "in", "not_in", "empty", "before_today", "not_contains"}

    def __init__(self, data: dict | None = None, schema: Schema = SCHEMA):
        data = data or _load("checks.yaml")
        self.checks = data["checks"]
        self.schema = schema
        self._validate()

    def _validate(self) -> None:
        seen = set()
        for k in self.checks:
            if k["id"] in seen:
                raise ValueError(f"Duplicate check id {k['id']}")
            seen.add(k["id"])
            if k["severity"] not in SEVERITY_ORDER:
                raise ValueError(f"Check {k['id']}: unknown severity {k['severity']}")
            for f in ("title", "action", "ref", "when"):
                if not k.get(f):
                    raise ValueError(f"Check {k['id']} is missing '{f}'")
            for cond in k["when"]:
                if "field" in cond:
                    if cond["field"] not in self.schema.ids:
                        raise ValueError(f"Check {k['id']}: unknown field {cond['field']}")
                    ops = set(cond) - {"field"}
                    if len(ops) != 1 or not ops <= self.OPS:
                        raise ValueError(f"Check {k['id']}: bad condition {cond}")
                elif not ({"expected_intensity", "intensity_mismatch"} & set(cond)):
                    raise ValueError(f"Check {k['id']}: bad condition {cond}")

    def _holds(self, cond: dict, row: Row, today: date) -> bool:
        if "expected_intensity" in cond:
            return expected_intensity(row, self.schema) in cond["expected_intensity"]
        if "intensity_mismatch" in cond:
            exp = expected_intensity(row, self.schema)
            got = row.get("governance_intensity")
            return bool(exp and got and exp != got) == bool(cond["intensity_mismatch"])
        v = row.get(cond["field"], "")
        if "eq" in cond:
            return v == str(cond["eq"])
        if "in" in cond:
            return v in [str(x) for x in cond["in"]]
        if "not_in" in cond:
            return bool(v) and v not in [str(x) for x in cond["not_in"]]
        if "empty" in cond:
            return (v == "") == bool(cond["empty"])
        if "not_contains" in cond:
            return cond["not_contains"].lower() not in v.lower()
        if "before_today" in cond:
            try:
                return date.fromisoformat(v) < today
            except ValueError:
                return False
        return False

    def run(self, rows: list[Row], today: date | None = None) -> list[Issue]:
        today = today or date.today()
        issues = []
        for r in rows:
            for k in self.checks:
                if all(self._holds(c, r, today) for c in k["when"]):
                    issues.append(Issue(r.get("use_case_id", ""), r.get("name", ""), k["id"], k["severity"],
                                        k["title"], k["action"], k["ref"]))
        issues.sort(key=lambda i: (SEVERITY_ORDER[i.severity], i.use_case_id))
        return issues


# --------------------------------------------------------------------------- #
# Import from the EU AI Act Risk Check
# --------------------------------------------------------------------------- #

RISK_CHECK_TIERS = {
    "minimal risk": "minimal",
    "transparency obligations": "limited",
    "high risk": "high",
    "prohibited": "prohibited",
    "out of scope": "out_of_scope",
}


def from_risk_check(markdown: str) -> dict[str, str]:
    """Read a Markdown report downloaded from the EU AI Act Risk Check.

    Returns the fields it can fill: name, eu_ai_act_tier, our_role, notes.
    """
    out: dict[str, str] = {}
    m = re.search(r"^#\s*EU AI Act risk check:\s*(.+)$", markdown, re.M | re.I)
    if m and m.group(1).strip().lower() != "unnamed ai system":
        out["name"] = m.group(1).strip()
    m = re.search(r"\*\*Result:\*\*\s*(.+)", markdown)
    if not m or m.group(1).strip().lower() not in RISK_CHECK_TIERS:
        raise ValueError("This does not look like an EU AI Act Risk Check report (no Result line found).")
    out["eu_ai_act_tier"] = RISK_CHECK_TIERS[m.group(1).strip().lower()]
    m = re.search(r"\*\*Role assessed:\*\*\s*(.+)", markdown)
    if m and m.group(1).strip().lower() in ("deployer", "provider"):
        out["our_role"] = m.group(1).strip().lower()
    d = re.search(r"\*\*Date:\*\*\s*(\d{4}-\d{2}-\d{2})", markdown)
    cites = re.findall(r"^- \*\*(Art[^*]+|Annex[^*]+)\.\*\*", markdown, re.M)
    note = "Classified with EU AI Act Risk Check" + (f" on {d.group(1)}" if d else "")
    if cites:
        note += f" ({'; '.join(c.strip() for c in cites[:3])})"
    out["notes"] = note
    return out


# --------------------------------------------------------------------------- #
# Summary and council report
# --------------------------------------------------------------------------- #


def summary(rows: list[Row]) -> dict[str, dict[str, int]]:
    out: dict[str, dict[str, int]] = {}
    for key in ("status", "eu_ai_act_tier", "governance_intensity"):
        counts: dict[str, int] = {}
        for r in rows:
            v = r.get(key) or "not set"
            counts[v] = counts.get(v, 0) + 1
        out[key] = counts
    return out


def council_report(rows: list[Row], issues: list[Issue], org: str = "", today: date | None = None) -> str:
    today = today or date.today()
    active = [r for r in rows if r.get("status") != "retired"]
    s = summary(active)
    high = [i for i in issues if i.severity == "high"]
    lines = [
        f"# AI register review{': ' + org if org else ''}",
        "",
        f"*Prepared {today.isoformat()} for the AI Council's quarterly register review.*",
        "",
        "## At a glance",
        "",
        f"- **Active use cases:** {len(active)} (live: {s['status'].get('live', 0)})",
        f"- **High-risk use cases:** {s['eu_ai_act_tier'].get('high', 0)}. Deployer duties for Annex III high-risk systems apply from 2 December 2027.",
        f"- **Open issues:** {len(issues)} ({len(high)} high)",
        "",
        "## By governance intensity",
        "",
        "| Intensity | Use cases |",
        "|---|---|",
    ]
    for k in ["stop", "redesign", "strict", "enhanced", "standard", "light", "not set"]:
        if s["governance_intensity"].get(k):
            lines.append(f"| {k} | {s['governance_intensity'][k]} |")
    lines += ["", "## Issues for decision", ""]
    if not issues:
        lines.append("No open issues.")
    else:
        lines += ["| Use case | Severity | Issue | Proposed action |", "|---|---|---|---|"]
        for i in issues:
            lines.append(f"| {i.use_case_id} {i.name} | {i.severity} | {i.title} | {i.action} |")
    upcoming = sorted((r for r in active if r.get("next_review") and r["next_review"] >= today.isoformat()),
                      key=lambda r: r["next_review"])[:5]
    if upcoming:
        lines += ["", "## Next reviews", "", "| Use case | Review due |", "|---|---|"]
        lines += [f"| {r['use_case_id']} {r['name']} | {r['next_review']} |" for r in upcoming]
    lines += ["", "---", "", "*Generated by the AI Use Case Register. Decisions are recorded by the AI Council, not by the tool.*", ""]
    return "\n".join(lines)
